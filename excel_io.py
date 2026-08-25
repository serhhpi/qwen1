# -*- coding: utf-8 -*-
"""
Всё про Excel: чтение присланных файлов в текст для промпта
и сборка итогового .xlsx из JSON-ответа Qwen.
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_ROWS_PER_SHEET = 400   # сколько строк листа отправляем модели
MAX_TABLE_CHARS = 60_000   # суммарный лимит текста таблиц в промпте


def read_workbooks(paths: list[Path]) -> list[dict]:
    """
    Читает каждый файл со всеми листами (заголовки + строки).
    Возвращает [{"file": ..., "sheets": [{"name", "headers", "rows"}]}]
    """
    workbooks = []
    for path in paths:
        try:
            sheets = pd.read_excel(path, sheet_name=None, dtype=str)
        except Exception as exc:
            raise ValueError(f"Не удалось прочитать {path.name}: {exc}") from exc

        wb = {"file": path.name, "sheets": []}
        for name, df in sheets.items():
            df = df.dropna(how="all").fillna("")
            wb["sheets"].append(
                {
                    "name": str(name),
                    "headers": [str(c) for c in df.columns],
                    "rows": [
                        [str(v) for v in row]
                        for row in df.head(MAX_ROWS_PER_SHEET).values.tolist()
                    ],
                }
            )
        workbooks.append(wb)
    return workbooks


def tables_to_text(workbooks: list[dict]) -> str:
    """Превращает прочитанные таблицы в компактный текст для промпта."""
    blocks: list[str] = []
    budget = MAX_TABLE_CHARS
    for wb in workbooks:
        for sheet in wb["sheets"]:
            title = f"### Файл: {wb['file']} | Лист: {sheet['name']}"
            lines = [title, " | ".join(sheet["headers"])]
            lines += [" | ".join(row) for row in sheet["rows"]]
            block = "\n".join(lines)
            if budget - len(block) < 0:
                blocks.append(f"{title}\n... данные сокращены (лимит символов)")
                budget = 0
                break
            blocks.append(block)
            budget -= len(block)
        if budget <= 0:
            break
    return "\n\n".join(blocks) or "(в файлах нет данных)"


def build_result_xlsx(payload: dict, out_path: Path) -> Path:
    """
    Собирает .xlsx из JSON модели:
        {"sheets": [{"name": ..., "headers": [...], "rows": [[...], ...]}]}
    Заголовки оформляются, ширины столбцов подгоняются, первая строка
    закрепляется — файл приятно открывать руками.
    """
    wb = Workbook()
    wb.remove(wb.active)

    thin = Side(style="thin", color="3A4A61")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", start_color="17324F")
    head_font = Font(bold=True, color="F2F6FC")

    for sheet in payload.get("sheets", []):
        name = str(sheet.get("name", "Лист"))[:31] or "Лист"
        ws = wb.create_sheet(title=name)

        headers = sheet.get("headers") or []
        rows = sheet.get("rows") or []

        ws.append(headers)
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = head_font
            cell.fill = head_fill
            cell.border = border
            cell.alignment = Alignment(vertical="center")

        for row in rows:  # 32767 — предел длины ячейки Excel
            ws.append([str(v)[:32767] for v in row])

        for col in range(1, len(headers) + 1):
            widths = [
                len(str(ws.cell(row=r, column=col).value or ""))
                for r in range(1, min(ws.max_row, 100) + 1)
            ]
            ws.column_dimensions[get_column_letter(col)].width = min(max(widths) + 4, 50)

        ws.freeze_panes = "A2"

    if not wb.sheetnames:
        wb.create_sheet("Пусто")
    wb.save(out_path)
    return out_path
